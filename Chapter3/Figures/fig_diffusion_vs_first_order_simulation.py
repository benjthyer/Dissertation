"""
Forward-model simulations comparing a purely diffusion-controlled process
(second-order, Fechtig & Kalbitzer 1966) against a purely first-order
(multi-energy-state, single-jump/exponential-decay) process, each reduced
under BOTH mathematical frameworks -- InvertFrequencyFactor1() (as if
diffusion-limited) and InvertFirstOrder() (as if first-order) -- to see how
each looks under its "native" reduction and under the "wrong" one, and where
the two frameworks agree or diverge. Run across four heating schedules: the
real Rustenburg pyrite step-heat (pulled from Reduced01_RustenburgPyUncert.xlsx),
two monotonic ramps, and a set of isothermal step sequences.

Scenarios:
    hematite_mdd_sim            -- Simulate_MDD_diffusion(): a hematite
                                    grain-size distribution (10% 0.01um, 30%
                                    0.1um, 40% 1um, 20% 10um), all domains
                                    sharing hematite's Ea/D0 (Farley 2018).
    multistate_first_order_sim  -- SimulateFirstOrder(): three independent
                                    first-order populations (40% k0=1.3e5
                                    Ea=170 kJ/mol, 40% k0=6.1e7 Ea=210 kJ/mol,
                                    20% k0=6.5e-1 Ea=43 kJ/mol).
    hematite_single_domain_sim  -- sanity check: Simulate_MDD_diffusion()
                                    with a single 0.1 um hematite domain (no
                                    domain mixing -- the "textbook" case both
                                    reductions are derived to handle exactly).
    first_order_single_state_sim -- sanity check: SimulateFirstOrder() with a
                                    single population, k0=1.3e5, Ea=170 kJ/mol
                                    (no state mixing, same rationale).

Per scenario x schedule, five plots. Steps with cumulative fractional yield
> 0.98 are excluded from all of them -- too close to fully degassed for
either per-step inversion to be numerically meaningful:
    <schedule>_diffusivity_arrhenius.pdf -- ln(D/a^2) vs 10^4/T
    <schedule>_lambda_arrhenius.pdf      -- ln(lambda) vs 10^4/T
    <schedule>_diff_vs_lambda.pdf        -- ln(D/a^2) vs ln(lambda), colored
                                             by cumulative yield, w/ regression
    <schedule>_lambda_diff_ratio.pdf     -- ln(lambda)/ln(D/a^2) vs cumulative
                                             yield, colored by temperature
    <schedule>_dRdT_vs_yield.pdf         -- d/dT of the above ratio vs
                                             cumulative yield

For the two monotonic-ramp schedules only (a release spectrum against
temperature isn't meaningful for a schedule that revisits temperatures):
    <schedule>_release_spectrum.pdf -- per-step (incremental) He release
                                        fraction vs Temperature, with the
                                        (constant) step duration in the title.

For every run of >=3 consecutive same-temperature steps found in a schedule
(by construction, all 12 four-step groups of the isothermal_sequences
schedule; incidentally, also any such runs that happen to occur elsewhere,
e.g. Rustenburg's occasional back-to-back repeat steps):
    <schedule>_isothermal_seq##_<T>C.pdf -- F-F0 vs t-t0 (t0/F0 = time/yield
                                        at the start of the run) with a
                                        bias-free (through-origin) quadratic
                                        fit F-F0 = a(t-t0)^2 + b(t-t0), t in
                                        seconds. Title reports R² (uncentered,
                                        appropriate for a through-origin fit),
                                        temperature, a/b^2, and -1/(1-F0) --
                                        for a pure first-order process these
                                        last two are related by a simple
                                        factor of 2 (Taylor-expanding
                                        F-F0=(1-F0)(1-exp(-k(t-t0))) gives
                                        a/b^2 = -1/(2(1-F0)) exactly), so
                                        printing both side by side lets you
                                        eyeball whether a given isothermal
                                        run behaves as first-order.

Outputs -> Chapter3/Figures/generated/<scenario>/<schedule>_*.pdf
"""
import os
import sys

import numpy as np
import openpyxl
import matplotlib.pyplot as plt
import scipy.stats

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)
from diffusion_utils import (Simulate_MDD_diffusion, SimulateFirstOrder,
                              InvertFrequencyFactor1, InvertFirstOrder,
                              Hem_Ea, Hem_D0_a2)
from plot_style import apply_style
apply_style()

GENERATED_DIR = os.path.join(SCRIPT_DIR, 'generated')

MAX_CUMULATIVE_YIELD = 0.98  # exclude steps beyond this Fcum from every plot
                             # -- too close to fully degassed for the per-step
                             # inversions to be numerically meaningful (see
                             # InvertFirstOrder / InvertFrequencyFactor's
                             # log((1-Fp)/(1-F))-type terms)

MIN_ISOTHERMAL_GROUP_LEN = 3  # minimum consecutive same-T steps to bother
                              # fitting the isothermal quadratic to -- a
                              # 2-parameter fit through exactly 2 points has
                              # no residual d.o.f. and a trivial R²=1

Ctemps = np.array([100, 200, 300, 400, 500, 600, 700, 800])

# ---------------------------------------------------------------------------
# Scenario 1: hematite multi-domain (grain-size) diffusion
# ---------------------------------------------------------------------------
# Hem_D0_a2 (diffusion_utils.py) is fit for a 1-micron reference domain, i.e.
# it's already D0/a^2 -- recover the pure prefactor D0 (cm^2/s) so
# Simulate_MDD_diffusion can re-normalize it by each domain's own radius.
HEM_D0 = Hem_D0_a2 * (1e-4) ** 2
HEM_GRAIN_RADII_CM = np.array([0.01, 0.1, 1, 10]) * 1e-4  # 0.01/0.1/1/10 micron
HEM_GRAIN_VOLUME_FRACTIONS = np.array([0.10, 0.30, 0.40, 0.20])

# ---------------------------------------------------------------------------
# Scenario 2: multi-energy-state first-order (single-jump) kinetics
# ---------------------------------------------------------------------------
FO_K0S = np.array([1.3e5, 6.1e7, 6.5e-1])
FO_EAS = np.array([170e3, 210e3, 43e3])  # J/mol
FO_FRACTIONS = np.array([0.40, 0.40, 0.20])

# ---------------------------------------------------------------------------
# Scenarios 3 & 4: single-domain / single-state sanity checks -- no mixing,
# so both reductions should behave "cleanly" and any curvature/kinks seen in
# scenarios 1-2 can be attributed to the mixing itself, not some artifact of
# the reduction machinery.
# ---------------------------------------------------------------------------
HEM_SINGLE_RADIUS_CM = np.array([0.1]) * 1e-4  # 0.1 micron
HEM_SINGLE_FRACTION = np.array([1.0])

FO_SINGLE_K0 = np.array([1.3e5])
FO_SINGLE_EA = np.array([170e3])  # J/mol
FO_SINGLE_FRACTION = np.array([1.0])


def simulate_hematite_mdd(temps_C, times_min):
    return np.asarray(Simulate_MDD_diffusion(
        D0=HEM_D0, Ea=Hem_Ea, T=temps_C, t=times_min,
        a=HEM_GRAIN_RADII_CM, v=HEM_GRAIN_VOLUME_FRACTIONS), dtype=float)


def simulate_multistate_first_order(temps_C, times_min):
    return np.asarray(SimulateFirstOrder(
        k0s=FO_K0S, Eas=FO_EAS, T=temps_C, t=times_min,
        v=FO_FRACTIONS), dtype=float)


def simulate_hematite_single_domain(temps_C, times_min):
    return np.asarray(Simulate_MDD_diffusion(
        D0=HEM_D0, Ea=Hem_Ea, T=temps_C, t=times_min,
        a=HEM_SINGLE_RADIUS_CM, v=HEM_SINGLE_FRACTION), dtype=float)


def simulate_first_order_single_state(temps_C, times_min):
    return np.asarray(SimulateFirstOrder(
        k0s=FO_SINGLE_K0, Eas=FO_SINGLE_EA, T=temps_C, t=times_min,
        v=FO_SINGLE_FRACTION), dtype=float)


SCENARIOS = {
    'hematite_mdd_sim': simulate_hematite_mdd,
    'multistate_first_order_sim': simulate_multistate_first_order,
    'hematite_single_domain_sim': simulate_hematite_single_domain,
    'first_order_single_state_sim': simulate_first_order_single_state,
}


# ---------------------------------------------------------------------------
# Heating schedules
# ---------------------------------------------------------------------------

def rustenburg_schedule():
    """Real Rustenburg pyrite heating schedule (Temp degC, Time min), pulled
    from the Diff Cell steps of Reduced01_RustenburgPyUncert.xlsx's 'Batch 55'
    sheet -- so the synthetic domains/populations see exactly the same T-t
    history as the real pyrite step-heat."""
    path = os.path.join(SCRIPT_DIR, '..', 'Tables', 'Reduced01_RustenburgPyUncert.xlsx')
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Batch 55']
    temps, times, names = [], [], []
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=5).value in ('Diff Cell', 'DiffCell'):
            temps.append(ws.cell(row=r, column=3).value)
            times.append(ws.cell(row=r, column=7).value)
            names.append(ws.cell(row=r, column=1).value)
    # Room-temperature cold-blank ('CB') checks are genuinely part of the
    # schedule (kept in the simulation so accumulated diffusive time stays
    # faithful to the real experiment) but release ~0 signal at 25 degC --
    # plotting them in log space produces meaningless extreme outliers that
    # dominate the axes and skew the regression, the same reason real
    # low-signal rows are excluded from the plotable dataset elsewhere in
    # this dissertation. Excluded from `plotable`, not from the schedule.
    plotable = np.array([not n.startswith('CB') for n in names])
    return np.array(temps, dtype=float), np.array(times, dtype=float), plotable


def ramp_schedule(t_start, t_end, t_step, minutes_per_step):
    temps = np.arange(t_start, t_end + 1, t_step, dtype=float)
    times = np.full(temps.shape, float(minutes_per_step))
    return temps, times, np.ones(temps.shape, dtype=bool)


def isothermal_sequence_schedule():
    """Three tiers of isothermal step sequences -- 4 repeated steps at each
    of 4 temperatures per tier, with shorter hold times as temperature (and
    therefore reaction/diffusion rate) increases: 60 min/step at 200-275C,
    30 min/step at 300-375C, 15 min/step at 400-475C. Each 4-step group is
    an "isothermal sequence" that plot_isothermal_quadratic_fits() fits a
    F-F0 vs t-t0 quadratic to."""
    tiers = [
        (60.0, (200, 225, 250, 275)),
        (30.0, (300, 325, 350, 375)),
        (15.0, (400, 425, 450, 475)),
    ]
    temps, times = [], []
    for minutes_per_step, tier_temps in tiers:
        for T in tier_temps:
            temps.extend([float(T)] * 4)
            times.extend([minutes_per_step] * 4)
    temps = np.array(temps, dtype=float)
    times = np.array(times, dtype=float)
    return temps, times, np.ones(temps.shape, dtype=bool)


SCHEDULES = {
    'rustenburg_schedule': rustenburg_schedule,
    'ramp_200-800C_5C_steps': lambda: ramp_schedule(200, 800, 5, 3.5),
    'ramp_225-775C_50C_steps': lambda: ramp_schedule(225, 775, 50, 8.0),
    'isothermal_sequences': isothermal_sequence_schedule,
}


def _is_monotonic_ramp(temps_C, times_min):
    """True for a strictly-increasing, constant-step-duration schedule (the
    two synthetic ramps) -- false for the real Rustenburg schedule, which
    revisits temperatures and varies step duration, so a release-fraction-
    vs-temperature spectrum wouldn't be single-valued."""
    return bool(np.all(np.diff(temps_C) > 0) and np.allclose(times_min, times_min[0]))


def _add_temp_axis(ax1):
    ax2 = ax1.twiny()
    xlim = ax1.get_xlim()
    ax2.set_xticks(1e4 / (Ctemps + 273))
    ax2.set_xlim(xlim)
    ax2.set_xticklabels(Ctemps)
    ax2.set_xlabel('Temperature (°C)')
    return ax2


def simulate_and_reduce(simulate_fn, temps_C, times_min):
    Fcum = simulate_fn(temps_C, times_min)
    times_s = times_min * 60
    d_a2 = np.asarray(InvertFrequencyFactor1(times_s, Fcum), dtype=float)
    lam = np.asarray(InvertFirstOrder(times_s, Fcum), dtype=float)
    tk = 10000.0 / (temps_C + 273.0)
    return tk, Fcum, d_a2, lam


# ---------------------------------------------------------------------------
# Plots
# ---------------------------------------------------------------------------

def plot_diffusivity_arrhenius(tk, d_a2, plotable, fcum, label, outdir):
    valid = np.isfinite(d_a2) & (d_a2 > 0) & plotable & (fcum <= MAX_CUMULATIVE_YIELD)
    fig, ax1 = plt.subplots()
    ax1.plot(tk[valid], np.log(d_a2[valid]), marker='o', markersize=4,
              linestyle=':', color='tab:blue', alpha=0.7)
    _add_temp_axis(ax1)
    ax1.set_xlabel(r'$10^4/T$ (K)')
    ax1.set_ylabel(r'$\ln(D/a^2)$ (s$^{-1}$)')
    ax1.grid(alpha=0.3)
    ax1.set_title(f'{label}: Simulated Diffusivity\n(2nd-order reduction)')
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, f'{label}_diffusivity_arrhenius.pdf'),
                dpi=300, bbox_inches='tight')
    plt.close(fig)


def plot_lambda_arrhenius(tk, lam, plotable, fcum, label, outdir):
    valid = np.isfinite(lam) & (lam > 0) & plotable & (fcum <= MAX_CUMULATIVE_YIELD)
    fig, ax1 = plt.subplots()
    ax1.plot(tk[valid], np.log(lam[valid]), marker='o', markersize=4,
              linestyle=':', color='tab:red', alpha=0.7)
    _add_temp_axis(ax1)
    ax1.set_xlabel(r'$10^4/T$ (K)')
    ax1.set_ylabel(r'$\ln(\lambda)$ (s$^{-1}$)')
    ax1.grid(alpha=0.3)
    ax1.set_title(f'{label}: Simulated Rate Constant\n(1st-order reduction)')
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, f'{label}_lambda_arrhenius.pdf'),
                dpi=300, bbox_inches='tight')
    plt.close(fig)


def plot_diff_vs_lambda(d_a2, lam, fcum, plotable, label, outdir):
    valid = (np.isfinite(d_a2) & (d_a2 > 0) & np.isfinite(lam) & (lam > 0)
             & plotable & (fcum <= MAX_CUMULATIVE_YIELD))
    if valid.sum() < 2:
        print(f"  fewer than 2 valid paired points for '{label}' -- skipping comparison plot")
        return

    x = np.log(lam[valid])
    y = np.log(d_a2[valid])
    c = fcum[valid]

    fig, ax = plt.subplots()
    sc = ax.scatter(x, y, c=c, cmap='viridis', vmin=c.min(), vmax=c.max(),
                     s=40, edgecolor='k', linewidth=0.4, zorder=2)
    cb = fig.colorbar(sc, ax=ax)
    cb.set_label(r'Cumulative $^3$He Yield, $F_i$')

    fit = scipy.stats.linregress(x, y)
    xfit = np.linspace(x.min(), x.max(), 50)
    ax.plot(xfit, fit.intercept + fit.slope * xfit, color='black', linestyle='--',
            label=f'slope = {fit.slope:.3f}, R² = {fit.rvalue ** 2:.3f}')

    ax.set_xlabel(r'$\ln(\lambda)$ (s$^{-1}$)')
    ax.set_ylabel(r'$\ln(D/a^2)$ (s$^{-1}$)')
    ax.grid(alpha=0.3)
    ax.legend(fontsize=9)
    ax.set_title(f'{label}: Diffusivity vs. Rate Constant')
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, f'{label}_diff_vs_lambda.pdf'),
                dpi=300, bbox_inches='tight')
    plt.close(fig)


def _lambda_diff_log_ratio(d_a2, lam):
    """R = ln(lambda) / ln(D/a^2), the shared quantity behind both the ratio
    plot and its dR/dT derivative below."""
    ln_d = np.log(d_a2, out=np.full_like(d_a2, np.nan), where=d_a2 > 0)
    ln_l = np.log(lam, out=np.full_like(lam, np.nan), where=lam > 0)
    with np.errstate(divide='ignore', invalid='ignore'):
        return ln_l / ln_d


def plot_lambda_diff_ratio(d_a2, lam, fcum, temps_C, plotable, label, outdir):
    ratio = _lambda_diff_log_ratio(d_a2, lam)
    valid = np.isfinite(ratio) & plotable & (fcum <= MAX_CUMULATIVE_YIELD)
    if valid.sum() < 1:
        print(f"  no valid points for '{label}' -- skipping lambda/diffusivity ratio plot")
        return

    fig, ax = plt.subplots()
    sc = ax.scatter(fcum[valid], ratio[valid], c=temps_C[valid], cmap='viridis',
                     vmin=temps_C[valid].min(), vmax=temps_C[valid].max(),
                     s=40, edgecolor='k', linewidth=0.4)
    cb = fig.colorbar(sc, ax=ax)
    cb.set_label('Temperature (°C)')

    ax.set_xlabel(r'Cumulative $^3$He Yield, $F_i$')
    ax.set_ylabel(r'$\ln(\lambda) \, / \, \ln(D/a^2)$')
    ax.grid(alpha=0.3)
    ax.set_title(f'{label}: Rate Constant / Diffusivity Log-Ratio')
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, f'{label}_lambda_diff_ratio.pdf'),
                dpi=300, bbox_inches='tight')
    plt.close(fig)


def plot_dRdT_vs_yield(d_a2, lam, fcum, temps_C, plotable, label, outdir):
    """dR/dT vs. cumulative yield, R = ln(lambda)/ln(D/a^2). A single domain
    or energy state should trace a smoothly varying R -- a new domain/state
    turning on and starting to dominate release shows up as a kink or spike
    in dR/dT, so this is a candidate diagnostic for detecting multi-domain
    (or multi-energy-state) behavior directly from the two reductions of the
    same dataset, without needing to already know which framework is correct.

    dR/dT is estimated by a finite difference between successive *valid*
    points (schedule order), which for the non-monotonic Rustenburg schedule
    means occasional negative or very large |dT| where the temperature path
    doubles back -- kept as-is (not clipped) since a reversal is itself a
    real feature of that schedule, but plotted against each interval's
    midpoint cumulative yield since yield -- unlike temperature -- is always
    monotonic non-decreasing in step order."""
    ratio = _lambda_diff_log_ratio(d_a2, lam)
    valid = np.isfinite(ratio) & plotable & (fcum <= MAX_CUMULATIVE_YIELD)
    idx = np.where(valid)[0]
    if len(idx) < 2:
        print(f"  fewer than 2 valid points for '{label}' -- skipping dR/dT plot")
        return

    dT = np.diff(temps_C[idx])
    with np.errstate(divide='ignore', invalid='ignore'):
        dRdT = np.diff(ratio[idx]) / dT
    fcum_mid = (fcum[idx][:-1] + fcum[idx][1:]) / 2
    finite = np.isfinite(dRdT)

    fig, ax = plt.subplots()
    ax.plot(fcum_mid[finite], dRdT[finite], marker='o', markersize=4,
            linestyle='-', color='tab:purple', alpha=0.8)
    ax.axhline(0, color='black', linewidth=0.8, alpha=0.5)
    ax.set_xlabel(r'Cumulative $^3$He Yield, $F_i$')
    ax.set_ylabel(r'$dR/dT$, $R = \ln(\lambda) \, / \, \ln(D/a^2)$ (°C$^{-1}$)')
    ax.grid(alpha=0.3)
    ax.set_title(f'{label}: d(λ/D Log-Ratio)/dT vs. Cumulative Yield')
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, f'{label}_dRdT_vs_yield.pdf'),
                dpi=300, bbox_inches='tight')
    plt.close(fig)


def plot_release_spectrum(temps_C, times_min, fcum, plotable, label, outdir):
    if not _is_monotonic_ramp(temps_C, times_min):
        return
    step_release = np.diff(np.concatenate(([0.0], fcum)))
    valid = plotable & (fcum <= MAX_CUMULATIVE_YIELD)

    fig, ax = plt.subplots()
    ax.plot(temps_C[valid], step_release[valid], marker='o', markersize=4,
            linestyle='-', color='tab:green', alpha=0.8)
    ax.set_xlabel('Temperature (°C)')
    ax.set_ylabel(r'Incremental $^3$He Release Fraction')
    ax.grid(alpha=0.3)
    ax.set_title(f'{label}: Helium Release Spectrum\n({times_min[0]:.1f} min per step)')
    fig.tight_layout()
    fig.savefig(os.path.join(outdir, f'{label}_release_spectrum.pdf'),
                dpi=300, bbox_inches='tight')
    plt.close(fig)


def _find_constant_T_runs(temps_C, plotable, min_len=MIN_ISOTHERMAL_GROUP_LEN):
    """Maximal runs of >=min_len consecutive plotable steps at (numerically)
    the same temperature -- the isothermal sequences the quadratic fit below
    is applied to."""
    runs = []
    n = len(temps_C)
    i = 0
    while i < n:
        if not plotable[i]:
            i += 1
            continue
        j = i
        while j + 1 < n and plotable[j + 1] and np.isclose(temps_C[j + 1], temps_C[i]):
            j += 1
        if j - i + 1 >= min_len:
            runs.append((i, j))
        i = j + 1
    return runs


def _fit_isothermal_quadratic(t_rel, F_rel):
    """Least-squares fit of F_rel = a*t_rel^2 + b*t_rel -- no intercept, since
    by construction F_rel=0 at t_rel=0 and a bias term would be physically
    meaningless for a curve anchored at its own origin. R² uses the
    "uncentered" convention appropriate for a through-origin fit: SS_tot is
    sum(F_rel^2) (the sum of squares of the F_rel=0 null model), not the
    usual sum((F_rel - mean)^2)."""
    X = np.column_stack([t_rel ** 2, t_rel])
    (a, b), *_ = np.linalg.lstsq(X, F_rel, rcond=None)
    resid = F_rel - (a * t_rel ** 2 + b * t_rel)
    ss_res = np.sum(resid ** 2)
    ss_tot = np.sum(F_rel ** 2)
    r2 = 1 - ss_res / ss_tot if ss_tot > 0 else np.nan
    return a, b, r2


def plot_isothermal_quadratic_fits(temps_C, times_min, fcum, plotable, label, outdir):
    runs = _find_constant_T_runs(temps_C, plotable)
    if not runs:
        return
    cumulative_seconds = np.cumsum(times_min) * 60.0

    for seq_i, (start, end) in enumerate(runs):
        t0 = cumulative_seconds[start - 1] if start > 0 else 0.0
        F0 = fcum[start - 1] if start > 0 else 0.0
        if F0 >= MAX_CUMULATIVE_YIELD:
            continue

        keep = fcum[start:end + 1] <= MAX_CUMULATIVE_YIELD
        if keep.sum() < MIN_ISOTHERMAL_GROUP_LEN:
            continue
        t_rel = cumulative_seconds[start:end + 1][keep] - t0
        F_rel = fcum[start:end + 1][keep] - F0

        a, b, r2 = _fit_isothermal_quadratic(t_rel, F_rel)
        T = temps_C[start]
        ab2 = a / b ** 2 if b != 0 else np.nan
        ref = -1.0 / (1.0 - F0)

        fig, ax = plt.subplots()
        ax.scatter(t_rel, F_rel, color='tab:orange', edgecolor='k', linewidth=0.4, zorder=2)
        tfit = np.linspace(0, t_rel.max(), 100)
        ax.plot(tfit, a * tfit ** 2 + b * tfit, color='black', linestyle='--', zorder=1)
        ax.set_xlabel(r'$t - t_0$ (s)')
        ax.set_ylabel(r'$F - F_0$')
        ax.grid(alpha=0.3)
        ax.set_title(
            f'{label}: Isothermal Fit @ {T:.0f}°C\n'
            f'R² = {r2:.4f},  a/b² = {ab2:.4g},  -1/(1-F₀) = {ref:.4g}'
        )
        fig.tight_layout()
        # F-F0 is often tiny (e.g. ~1e-8 at low T for a slow process), which
        # makes matplotlib draw a "x1eN" scientific-notation offset above the
        # y-axis -- reserve extra headroom so it doesn't collide with the
        # two-line title above.
        fig.subplots_adjust(top=0.82)
        fname = f'{label}_isothermal_seq{seq_i:02d}_{T:.0f}C.pdf'
        fig.savefig(os.path.join(outdir, fname), dpi=300, bbox_inches='tight')
        plt.close(fig)


def main():
    for scenario_name, simulate_fn in SCENARIOS.items():
        outdir = os.path.join(GENERATED_DIR, scenario_name)
        os.makedirs(outdir, exist_ok=True)
        for label, schedule_fn in SCHEDULES.items():
            print(f"Simulating '{scenario_name}' / '{label}'...")
            temps_C, times_min, plotable = schedule_fn()
            tk, Fcum, d_a2, lam = simulate_and_reduce(simulate_fn, temps_C, times_min)
            plot_diffusivity_arrhenius(tk, d_a2, plotable, Fcum, label, outdir)
            plot_lambda_arrhenius(tk, lam, plotable, Fcum, label, outdir)
            plot_diff_vs_lambda(d_a2, lam, Fcum, plotable, label, outdir)
            plot_lambda_diff_ratio(d_a2, lam, Fcum, temps_C, plotable, label, outdir)
            plot_dRdT_vs_yield(d_a2, lam, Fcum, temps_C, plotable, label, outdir)
            plot_release_spectrum(temps_C, times_min, Fcum, plotable, label, outdir)
            plot_isothermal_quadratic_fits(temps_C, times_min, Fcum, plotable, label, outdir)
        print(f"Wrote figures to {outdir}")


if __name__ == '__main__':
    main()
