"""Rebuild the plot-ready sulfide apparent-diffusivity / apparent-rate-constant cache.

Diffusivity data (Fim, 10000/TK, ln(d/a^2), and its LB/UB bound widths) is pulled
from the "Diffusivities" (or "Sheet2") tab of each Reduced01_*.xlsx workbook.
Bounded apparent rate constants (FLB/Fm/FUB, lambdaLB/lambdaM/lambdaUB) are pulled
from the tail of the first tab of the same workbook.

Not every row in a Reduced01 workbook is plotable -- some are standards, blanks,
permeation steps, or have unconstrained uncertainties, and are flagged by row
color in the original workbooks. Rather than re-deriving plotability from color,
this script uses the existing plotable_sulfides.xlsx as a one-time key: it is
matched against the *current* Reduced01 files on 10000/TK (a temperature-derived
quantity that a diffusivity-bounds methodology change won't touch) to recover
which Step numbers are plotable for each sample. That key is persisted to
plotable_row_index.json so later re-runs -- after the diffusivity/lambda values
themselves change under a further methodology update -- no longer need
plotable_sulfides.xlsx at all.

Usage:
    python build_plotable_cache.py [--rebuild-index]

Output:
    plotable_row_index.json   -- {sample: {source_file, steps: [...]}}
    plotable_sulfides_cache.pkl -- tidy pandas DataFrame, one row per plotable step
"""

import argparse
import json
from pathlib import Path

import openpyxl
import pandas as pd

TABLES_DIR = Path(__file__).resolve().parent
PLOTABLE_XLSX = TABLES_DIR / "plotable_sulfides.xlsx"
INDEX_JSON = TABLES_DIR / "plotable_row_index.json"
CACHE_PKL = TABLES_DIR / "plotable_sulfides_cache.pkl"

TK_TOL = 1e-6  # relative tolerance matching 10000/TK between the old cache and current files

LAMBDA_NAMES = ("FLB", "Fm", "FUB", "λLB", "λm", "λUB")
LAMBDA_OUT_NAMES = ("FLB", "Fm", "FUB", "lambdaLB", "lambdaM", "lambdaUB")


def _reduced_files():
    return sorted(TABLES_DIR.glob("Reduced01_*.xlsx"))


def _diffusivity_sheet(wb):
    return wb["Diffusivities"] if "Diffusivities" in wb.sheetnames else wb["Sheet2"]


def _diff_columns(ws):
    """Locate columns of interest by header text rather than fixed index, since
    the exact column count varies a little between samples."""
    row1 = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    row3 = [ws.cell(row=3, column=c).value for c in range(1, ws.max_column + 1)]
    delta_idx = [i + 1 for i, h in enumerate(row1) if h == "Δln(d/a^2)"]
    lnda2_idx = [i + 1 for i, h in enumerate(row1) if h == "ln(d/a^2)"]
    return dict(
        step=row3.index("Step") + 1,
        tk=row1.index("10000/TK") + 1,
        fim=row1.index("Fim") + 1,
        lnda2=lnda2_idx[0],  # first occurrence = the central ("m") estimate
        lb=delta_idx[0],
        ub=delta_idx[1],
    )


def _diff_data_rows(ws, cols):
    return [
        r for r in range(4, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=cols["tk"]).value, (int, float))
    ]


# Both the first sheet and Diffusivities/Sheet2 are built from the same
# ordered list of step-heat analyses, offset only by their fixed header-row
# counts (first sheet: 1 header row, data from row 2; Diffusivities: 3 header
# rows, data from row 4) -- so row *position* is the reliable join key.
# Row *values* (Step / He RefNo) are not: Rustenburg's first sheet is missing
# He RefNo 90582 entirely (a step excised from that sheet), so Diffusivities'
# own Step column -- a naive incrementing counter, not read back from the
# first sheet -- silently drifts one off from the true He RefNo for every
# row after the gap. Verified by cross-checking Temp agreement at this fixed
# offset across every data row of all 7 Reduced01 files (including
# Rustenburg's gap), with zero mismatches, vs. 28 mismatches when joining on
# Step/He RefNo value.
WS1_DATA_START_ROW = 2
DIFF_DATA_START_ROW = 4
ROW_OFFSET = DIFF_DATA_START_ROW - WS1_DATA_START_ROW


def _first_sheet_row(r2):
    """First-sheet row that is the same physical step as Diffusivities row r2."""
    return r2 - ROW_OFFSET


def _lambda_columns(ws1):
    header = [ws1.cell(row=1, column=c).value for c in range(1, ws1.max_column + 1)]
    if not all(n in header for n in LAMBDA_NAMES):
        return None
    return {n: header.index(n) + 1 for n in LAMBDA_NAMES}


def _sequential_match(cache_rows, ws2, cols):
    """Walk the old cache rows and the current Diffusivities data rows in
    lockstep, matching on 10000/TK. Returns None (no match) as soon as any row
    fails to find its next occurrence in order -- callers try this against
    every Reduced01 file and require exactly one to succeed."""
    data_rows = _diff_data_rows(ws2, cols)
    pointer = 0
    matched = []
    for _fim, tk, *_rest in cache_rows:
        found = None
        for j in range(pointer, len(data_rows)):
            r = data_rows[j]
            v = ws2.cell(row=r, column=cols["tk"]).value
            if v is not None and abs(v - tk) < TK_TOL * max(abs(tk), 1):
                found = (r, j)
                break
        if found is None:
            return None
        matched.append(found[0])
        pointer = found[1] + 1
    return matched


def build_row_index():
    cache_wb = openpyxl.load_workbook(PLOTABLE_XLSX, data_only=True)
    reduced_wbs = {f: openpyxl.load_workbook(f, data_only=True) for f in _reduced_files()}

    index = {}
    for sample in cache_wb.sheetnames:
        cs = cache_wb[sample]
        cache_rows = [row for row in cs.iter_rows(min_row=2, values_only=True) if row[1] is not None]
        if not cache_rows:
            continue

        candidates = []
        for f, wb in reduced_wbs.items():
            ws2 = _diffusivity_sheet(wb)
            cols = _diff_columns(ws2)
            matched = _sequential_match(cache_rows, ws2, cols)
            if matched is not None:
                candidates.append((f, matched, cols))

        if len(candidates) != 1:
            names = [c[0].name for c in candidates]
            raise RuntimeError(
                f"Sample '{sample}': expected exactly one matching Reduced01 file, "
                f"found {len(candidates)}: {names}"
            )

        f, matched_rows, cols = candidates[0]
        ws2 = _diffusivity_sheet(reduced_wbs[f])
        steps = [ws2.cell(row=r, column=cols["step"]).value for r in matched_rows]
        index[sample] = {"source_file": f.name, "steps": steps}
        print(f"{sample}: {len(steps)} plotable rows -> {f.name}")

    INDEX_JSON.write_text(json.dumps(index, indent=2))
    print(f"Wrote row-index key to {INDEX_JSON}")
    return index


def build_cache(index):
    records = []
    warned_samples = set()

    for sample, entry in index.items():
        f = TABLES_DIR / entry["source_file"]
        wb = openpyxl.load_workbook(f, data_only=True)

        ws2 = _diffusivity_sheet(wb)
        cols = _diff_columns(ws2)
        step_to_row2 = {
            ws2.cell(row=r, column=cols["step"]).value: r
            for r in _diff_data_rows(ws2, cols)
        }

        ws1 = wb[wb.sheetnames[0]]
        lam_cols = _lambda_columns(ws1)
        if lam_cols is None and sample not in warned_samples:
            print(
                f"WARNING: {entry['source_file']} has no {'/'.join(LAMBDA_NAMES)} "
                f"block yet -- lambda columns will be NaN for sample '{sample}'."
            )
            warned_samples.add(sample)

        for step in entry["steps"]:
            r2 = step_to_row2.get(step)
            if r2 is None:
                print(f"WARNING: step {step} for sample '{sample}' no longer exists in "
                      f"{entry['source_file']} -- skipping row.")
                continue

            rec = {
                "sample": sample,
                "step": step,
                "Fim": ws2.cell(row=r2, column=cols["fim"]).value,
                "10000/TK": ws2.cell(row=r2, column=cols["tk"]).value,
                "ln(d/a^2)": ws2.cell(row=r2, column=cols["lnda2"]).value,
                "LB_dln(d/a^2)": ws2.cell(row=r2, column=cols["lb"]).value,
                "UB_dln(d/a^2)": ws2.cell(row=r2, column=cols["ub"]).value,
            }

            r1 = _first_sheet_row(r2) if lam_cols is not None else None
            if r1 is not None and not (WS1_DATA_START_ROW <= r1 <= ws1.max_row):
                print(f"WARNING: {sample} step {step}: positionally-joined first-sheet "
                      f"row {r1} is out of range -- lambda columns will be NaN for this row.")
                r1 = None
            for src_name, out_name in zip(LAMBDA_NAMES, LAMBDA_OUT_NAMES):
                rec[out_name] = ws1.cell(row=r1, column=lam_cols[src_name]).value if r1 else None

            records.append(rec)

    df = pd.DataFrame.from_records(records)

    for col in df.columns:
        is_error = df[col].apply(lambda v: isinstance(v, str) and v.startswith("#"))
        if is_error.any():
            for _, row in df.loc[is_error, ["sample", "step", col]].iterrows():
                print(f"WARNING: {row['sample']} step {row['step']}: {col} = {row[col]} "
                      f"(Excel formula error carried through as-is, not converted to NaN)")

    df.to_pickle(CACHE_PKL)
    print(f"Wrote {len(df)} rows across {df['sample'].nunique()} samples to {CACHE_PKL}")
    return df


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--rebuild-index",
        action="store_true",
        help="Re-derive plotable_row_index.json from plotable_sulfides.xlsx even if it already exists.",
    )
    args = parser.parse_args()

    if args.rebuild_index or not INDEX_JSON.exists():
        index = build_row_index()
    else:
        index = json.loads(INDEX_JSON.read_text())

    build_cache(index)


if __name__ == "__main__":
    main()
